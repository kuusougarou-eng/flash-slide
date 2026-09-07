# -*- coding: utf-8 -*-
"""
参照スライド解析の E2E テスト用に、デザインの異なる 3 種のデッキを生成する。
  python test/fixtures/make-decks.py
出力: test/fixtures/deck-A-placeholder.pptx / deck-B-copypaste.pptx / deck-C-darkband.pptx

A: 標準テーマのプレースホルダ(タイトル + コンテンツ)。リード文なし。フッター/ページ番号あり。
B: コピペ運用デッキ。テキストボックスのみ。ネイビー 24pt タイトル + グレー 16pt リード + 罫線 + 出典 + ページ番号。
C: 上部にダークバンド、白抜きタイトル、リード無し、右下にロゴ風図形、4:3 サイズ。
"""
import os, sys
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUT = os.path.dirname(os.path.abspath(__file__))

def pt(v):
    return Emu(int(v * 12700))

def add_text(slide, x, y, w, h, text, size, bold=False, color=RGBColor(0x1A, 0x1A, 0x1A), align=PP_ALIGN.LEFT, font=None, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(pt(x), pt(y), pt(w), pt(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    if font:
        r.font.name = font
    return tb

def add_rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, pt(x), pt(y), pt(w), pt(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    return s

# ---------- A: placeholder deck ----------
def deck_a():
    prs = Presentation()  # 既定 4:3 → 16:9 に
    prs.slide_width = pt(960)
    prs.slide_height = pt(540)
    layout = prs.slide_layouts[1]  # Title and Content
    for i in range(2):
        s = prs.slides.add_slide(layout)
        s.shapes.title.text = f"既存スライド {i+1}: 事業環境の整理"
        body = s.placeholders[1]
        body.text_frame.text = "国内市場は年率 4% で縮小"
        body.text_frame.add_paragraph().text = "上位 3 社の寡占が進行"
        add_text(s, 40, 505, 300, 24, "Company Confidential", 10, color=RGBColor(0x88, 0x88, 0x88))
        add_text(s, 880, 505, 60, 24, str(i + 1), 10, color=RGBColor(0x88, 0x88, 0x88), align=PP_ALIGN.RIGHT)
    prs.save(os.path.join(OUT, "deck-A-placeholder.pptx"))

# ---------- B: copy-paste textbox deck ----------
def deck_b():
    prs = Presentation()
    prs.slide_width = pt(960)
    prs.slide_height = pt(540)
    blank = prs.slide_layouts[6]
    navy = RGBColor(0x1F, 0x3A, 0x5F)
    for i in range(2):
        s = prs.slides.add_slide(blank)
        add_text(s, 48, 28, 864, 44, f"既存スライド {i+1}: 主力製品の粗利率は 3 年で 6pt 低下", 24, bold=True, color=navy, font="Meiryo UI", anchor=MSO_ANCHOR.MIDDLE)
        add_rect(s, 48, 76, 864, 1.5, navy)
        add_text(s, 48, 84, 864, 40, "原材料高と値引き競争が主因。価格改定と SKU 整理で来期 2pt の回復を狙う。", 16, color=RGBColor(0x59, 0x59, 0x59), font="Meiryo UI")
        # 本文(前スライドの中身。追加時には消される想定)
        add_rect(s, 48, 140, 270, 300, RGBColor(0xF2, 0xF2, 0xF2))
        add_text(s, 60, 150, 246, 280, "本文ダミー: 粗利率推移のグラフ", 14, font="Meiryo UI")
        add_rect(s, 345, 140, 567, 300, RGBColor(0xF2, 0xF2, 0xF2))
        add_text(s, 357, 150, 543, 280, "本文ダミー: 要因分解と打ち手", 14, font="Meiryo UI")
        # フッター
        add_text(s, 48, 470, 600, 20, "出典: 社内管理会計データ(2026年3月期)", 10, color=RGBColor(0x88, 0x88, 0x88), font="Meiryo UI")
        add_rect(s, 48, 500, 864, 0.75, RGBColor(0xBF, 0xBF, 0xBF))
        add_text(s, 48, 506, 400, 20, "© 2026 Sample Consulting Inc.", 9, color=RGBColor(0x88, 0x88, 0x88), font="Meiryo UI")
        add_text(s, 880, 506, 32, 20, str(i + 1), 9, color=RGBColor(0x88, 0x88, 0x88), align=PP_ALIGN.RIGHT, font="Meiryo UI")
    prs.save(os.path.join(OUT, "deck-B-copypaste.pptx"))

# ---------- C: dark band deck (4:3) ----------
def deck_c():
    prs = Presentation()
    prs.slide_width = pt(720)
    prs.slide_height = pt(540)
    blank = prs.slide_layouts[6]
    dark = RGBColor(0x2B, 0x2B, 0x2B)
    green = RGBColor(0x0B, 0x7A, 0x4B)
    for i in range(2):
        s = prs.slides.add_slide(blank)
        add_rect(s, 0, 0, 720, 64, dark)
        add_rect(s, 0, 64, 720, 4, green)
        add_text(s, 30, 8, 660, 48, f"既存スライド {i+1}: 施策ロードマップ", 22, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, 30, 90, 660, 380, "本文ダミー(箇条書き)\n・施策 1\n・施策 2", 16)
        # ロゴ風(右下の緑ブロック + 文字)
        add_rect(s, 620, 496, 70, 24, green)
        add_text(s, 620, 496, 70, 24, "LOGO", 10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, 30, 500, 300, 20, f"{i+1} / 2", 9, color=RGBColor(0x88, 0x88, 0x88))
    prs.save(os.path.join(OUT, "deck-C-darkband.pptx"))

deck_a(); deck_b(); deck_c()
print("decks written to", OUT)
